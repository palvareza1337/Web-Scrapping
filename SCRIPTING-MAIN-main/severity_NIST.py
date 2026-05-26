import requests
import json
import time
import random
from docx import Document

# Diccionario para almacenar las CVE encontradas con sus puntajes
found_cve = dict()

def get_cves_base_scores(cve_list):
    scores = list()
    for cve in cve_list:
        time.sleep(random.random())
        scores.append(get_cve_base_score(cve))
    return scores

def get_cve_base_score(cve):
    http_headers = {
        'User-Agent': 'Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/97.0.4692.71 Safari/537.36',
        'apiKey': '278330a8-0c9e-447c-9dc7-c1f1fc3ad9a8'
    }

    if cve in found_cve.keys():
        return (cve, found_cve.get(cve))
    
    url = f"https://services.nvd.nist.gov/rest/json/cves/2.0?cveId={cve}"
    try:
        response = requests.get(url, headers=http_headers)
        data = json.loads(response.text)
    
    except:
        time.sleep(10)

        try:
            response = requests.get(url, headers=http_headers)
            data = json.loads(response.text)
        except:
            data = ''
            print(f"Favor de revisar {cve} manualmente.")
            print(f"Favor de revisar {cve} manualmente.")
            print(f"Favor de revisar {cve} manualmente.")

    try:
        cve_value = data["vulnerabilities"][0]["cve"]["metrics"]["cvssMetricV31"][0]["cvssData"]["baseScore"]
    except:
        cve_value = 0
    
    found_cve[cve] = cve_value
    return (cve, cve_value)

from docx.shared import Inches

def export_cves_to_docx(cves_scores_by_ip, document):
    
    document.add_heading('CVEs y Puntajes Base por IP', 0)
    
    for ip, cves_scores in cves_scores_by_ip.items():
        document.add_heading(f'IP: {ip}', level=1)
        table = document.add_table(rows=1, cols=2)
        table.style = 'Table Grid'
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'CVE'
        hdr_cells[1].text = 'Puntaje Base'

        for cve, score in cves_scores:
            row_cells = table.add_row().cells
            row_cells[0].text = cve
            row_cells[1].text = str(score)
        
        # Agregar enlaces de referencia al final de cada sección de IP
        document.add_paragraph('Referencias:')
        for cve, _ in cves_scores:
            link = f"https://nvd.nist.gov/vuln/detail/{cve}"
            paragraph = document.add_paragraph()
            run = paragraph.add_run(cve)
            run.hyperlink = link
            paragraph.add_run(f" - {link}")

    return document



if __name__ == "__main__":
    print(get_cve_base_score(""))
