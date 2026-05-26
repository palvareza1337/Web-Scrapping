from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, color):
    """
    Aplica un color de fondo (en hexadecimal, sin el símbolo #) a la celda.
    """
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), color)
    tcPr.append(shd)

# Crear el documento
doc = Document()

# Configuración de fuente y tamaño
fuente = "Calibri (Cuerpo)"
tamano = Pt(9)

# ===============================================================
# TABLA ÚNICA: CABECERA + CASO (8 filas, 2 columnas para cabecera y 1 columna para caso)
# ===============================================================
table = doc.add_table(rows=8, cols=2)
table.style = 'Table Grid'

merged_cell = table.rows[0].cells[0].merge(table.rows[0].cells[1])
merged_cell.text = ""
set_cell_background(merged_cell, "FF0000")

header_data = [
    ("NOMBRE", "Correos encontrados en filtraciones"),
    ("SEVERIDAD", "ALTA"),
    ("CVSS BASE", ""),
    ("CVE", ""),
    ("ACCESO", ""),
    ("RIESGO", "ALTO")
]

header_bg = "D9D9D9"
for i, (campo, valor) in enumerate(header_data, start=1):
    cell_left = table.rows[i].cells[0]
    para_left = cell_left.paragraphs[0]
    run_left = para_left.add_run(campo)
    run_left.font.size = tamano
    run_left.font.name = fuente
    set_cell_background(cell_left, header_bg)

    cell_right = table.rows[i].cells[1]
    if valor:
        para_right = cell_right.paragraphs[0]
        run_right = para_right.add_run(valor)
        run_right.font.size = tamano
        run_right.font.name = fuente

header_oscuro = "666666"
cell_caso_header = table.rows[7].cells[0].merge(table.rows[7].cells[1])
para_caso_header = cell_caso_header.paragraphs[0]
run_caso_header = para_caso_header.add_run("CASO")
run_caso_header.bold = True
run_caso_header.font.size = tamano
run_caso_header.font.name = fuente
run_caso_header.font.color.rgb = RGBColor(255, 255, 255)
set_cell_background(cell_caso_header, header_oscuro)

cell_caso = table.add_row().cells[0]
cell_caso.merge(table.rows[8].cells[1])
cell_caso.text = "En la siguiente lista se detallan las filtraciones identificadas."

# ===============================================================
# TABLA DE MITIGACIÓN
# ===============================================================
table_mitigacion = doc.add_table(rows=2, cols=1)
table_mitigacion.style = 'Table Grid'

cell_mitigacion_header = table_mitigacion.rows[0].cells[0]
para_mitigacion_header = cell_mitigacion_header.paragraphs[0]
run_mitigacion_header = para_mitigacion_header.add_run("MITIGACIÓN")
run_mitigacion_header.bold = True
run_mitigacion_header.font.size = tamano
run_mitigacion_header.font.name = fuente
run_mitigacion_header.font.color.rgb = RGBColor(255, 255, 255)
set_cell_background(cell_mitigacion_header, header_oscuro)

table_mitigacion.rows[1].cells[0].text = ("Se recomienda mantener una alerta con respecto a las cuentas mencionadas ya \n"
                                         "que podrían ser incluidas en alguna campaña masiva de phishing.")

# ===============================================================
# TABLA DE EVIDENCIA
# ===============================================================
table_evidencia = doc.add_table(rows=2, cols=1)
table_evidencia.style = 'Table Grid'

cell_evidencia_header = table_evidencia.rows[0].cells[0]
para_evidencia_header = cell_evidencia_header.paragraphs[0]
run_evidencia_header = para_evidencia_header.add_run("EVIDENCIA")
run_evidencia_header.bold = True
run_evidencia_header.font.size = tamano
run_evidencia_header.font.name = fuente
run_evidencia_header.font.color.rgb = RGBColor(255, 255, 255)
set_cell_background(cell_evidencia_header, header_oscuro)

# --- Contenido de EVIDENCIA ---
cell_evidencia = table_evidencia.rows[1].cells[0]
para_evidencia = cell_evidencia.paragraphs[0]
run_evidencia = para_evidencia.add_run("Se identificaron los siguientes correos afectados:\n")
run_evidencia.font.size = tamano
run_evidencia.font.name = fuente

# Leer y agregar correos desde el archivo reporte_brechas.txt
with open("reporte_brechas.txt", "r") as file:
    lines = file.readlines()

for line in lines:
    run_correo = para_evidencia.add_run(line.strip() + "\n")
    run_correo.font.size = tamano
    run_correo.font.name = fuente

# Guardar el documento resultante
doc.save("filtraciones_generado.docx")
print("Documento generado: filtraciones_generado.docx")
