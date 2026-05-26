from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), color)
    tcPr.append(shd)

# Crear documento
doc = Document()

fuente = "Calibri (Cuerpo)"
tamano = Pt(9)

# ===============================================================
# TABLA PRINCIPAL
# ===============================================================
# 0: barra azul
# 1-4: cabecera
# 5: headers descripción/impacto
# 6: contenido
table = doc.add_table(rows=7, cols=2)
table.style = 'Table Grid'

# --- Fila 0 (barra azul) ---
merged_cell = table.rows[0].cells[0].merge(table.rows[0].cells[1])
merged_cell.text = ""
set_cell_background(merged_cell, "0070C0")

# --- Cabecera ---
header_data = [
    ("NOMBRE", "Análisis Perimetral"),
    ("SEVERIDAD", ""),
    ("ACCESO", ""),
    ("RIESGO", "")
]

header_bg = "D9D9D9"

for i, (campo, valor) in enumerate(header_data, start=1):
    cell_left = table.rows[i].cells[0]
    p = cell_left.paragraphs[0]
    run = p.add_run(campo)
    run.font.size = tamano
    run.font.name = fuente
    set_cell_background(cell_left, header_bg)

    cell_right = table.rows[i].cells[1]
    if valor:
        p2 = cell_right.paragraphs[0]
        run2 = p2.add_run(valor)
        run2.font.size = tamano
        run2.font.name = fuente

# --- Headers DESCRIPCIÓN / IMPACTO ---
header_dark = "666666"

cell_desc = table.rows[5].cells[0]
p = cell_desc.paragraphs[0]
run = p.add_run("DESCRIPCIÓN")
run.bold = True
run.font.size = tamano
run.font.name = fuente
run.font.color.rgb = RGBColor(255, 255, 255)
set_cell_background(cell_desc, header_dark)

cell_imp = table.rows[5].cells[1]
p = cell_imp.paragraphs[0]
run = p.add_run("IMPACTO")
run.bold = True
run.font.size = tamano
run.font.name = fuente
run.font.color.rgb = RGBColor(255, 255, 255)
set_cell_background(cell_imp, header_dark)

# --- Contenido vacío (puedes llenarlo después) ---
for j in range(2):
    cell = table.rows[6].cells[j]
    cell.text = ""
    for p in cell.paragraphs:
        for run in p.runs:
            run.font.size = tamano
            run.font.name = fuente

# ===============================================================
# BLOQUES DE REGISTROS
# ===============================================================
secciones = [
    "REGISTROS DE DOMINIO",
    "REGISTROS NS",
    "REGISTROS TIPO A",
    "REGISTROS MX",
    "REGISTROS TXT"
]

for seccion in secciones:
    t = doc.add_table(rows=2, cols=1)
    t.style = 'Table Grid'

    # Header
    cell_header = t.rows[0].cells[0]
    p = cell_header.paragraphs[0]
    run = p.add_run(seccion)
    run.bold = True
    run.font.size = tamano
    run.font.name = fuente
    run.font.color.rgb = RGBColor(255, 255, 255)
    set_cell_background(cell_header, header_dark)

    # Contenido vacío
    cell_body = t.rows[1].cells[0]
    cell_body.text = ""
    for p in cell_body.paragraphs:
        for run in p.runs:
            run.font.size = tamano
            run.font.name = fuente

# Guardar
doc.save("analisis_perimetral_generado.docx")
print("Documento generado: analisis_perimetral_generado.docx")