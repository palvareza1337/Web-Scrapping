from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def agregar_tabla_vulnerabilidades(document, ip, cves):

    def set_bg(cell, color):
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), color)
        tcPr.append(shd)

    def fmt(cell, text):
        cell.text = text
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.size = Pt(9)
                run.font.name = "Calibri (Cuerpo)"

    # CVEs
    lista_cves = []

    for cve in cves:
        if isinstance(cve, dict):
            lista_cves.append(cve.get("id", ""))
        elif isinstance(cve, tuple):
            lista_cves.append(cve[0])  # ← aquí está la clave
        elif isinstance(cve, str):
            lista_cves.append(cve)

    evidencia = "\n".join(lista_cves)

    fuentes = "\n".join(
        [f"https://nvd.nist.gov/vuln/detail/{cve}" for cve in lista_cves]
    )

    # Tabla
    table = document.add_table(rows=11, cols=2)
    table.style = 'Table Grid'

    # Barra roja
    merged = table.rows[0].cells[0].merge(table.rows[0].cells[1])
    set_bg(merged, "C00000")

    # Cabecera
    header_data = [
        ("NOMBRE", f"Servidor vulnerable ({ip})"),
        ("SEVERIDAD", "ALTA"),
        ("CVSS BASE", ""),
        ("CVE", ", ".join(lista_cves[:5]) + ("..." if len(lista_cves) > 5 else "")),
        ("ACCESO", ip),
        ("RIESGO", "Exposición de múltiples vulnerabilidades conocidas")
    ]

    for i, (k, v) in enumerate(header_data, start=1):
        set_bg(table.rows[i].cells[0], "D9D9D9")
        fmt(table.rows[i].cells[0], k)
        fmt(table.rows[i].cells[1], v)

    # 🔹 DESCRIPCIÓN / IMPACTO
    headers = ["DESCRIPCIÓN", "IMPACTO"]
    for i, txt in enumerate(headers):
        cell = table.rows[7].cells[i]
        set_bg(cell, "666666")
        p = cell.paragraphs[0]
        run = p.add_run(txt)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(255, 255, 255)

    fmt(table.rows[8].cells[0], "Se identificaron múltiples vulnerabilidades públicas asociadas a este activo.")
    fmt(table.rows[8].cells[1], "Un atacante podría explotar estas vulnerabilidades dependiendo del contexto del servicio.")

    # 🔹 EVIDENCIA / FUENTES
    headers2 = ["EVIDENCIA", "FUENTES"]
    for i, txt in enumerate(headers2):
        cell = table.rows[9].cells[i]
        set_bg(cell, "666666")
        p = cell.paragraphs[0]
        run = p.add_run(txt)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(255, 255, 255)

    fmt(table.rows[10].cells[0], evidencia)
    fmt(table.rows[10].cells[1], fuentes)