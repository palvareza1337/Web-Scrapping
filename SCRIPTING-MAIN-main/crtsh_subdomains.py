# crtsh_subdomains.py
import requests
import pandas as pd
import socket
import time
from bs4 import BeautifulSoup
from docx.enum.text import WD_ALIGN_PARAGRAPH

# ------- CONFIGURACIÓN -------
USER_AGENT = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
JSON_TIMEOUT = (8, 12)   # (connect_timeout, read_timeout)
HTML_TIMEOUT = (8, 20)
MAX_RETRIES = 2
REQUESTS_SLEEP = 2
# -----------------------------

def fetch_data(domain, verbose=True):
    start_total = time.time()
    if verbose: print(f"[+] Obteniendo subdominios desde crt.sh para: {domain}")

    subdomains = {}  # dict: {subdomain: "Sí"/"No"}
    json_success = False

    # --- Intento JSON ---
    url_json = f"https://crt.sh/json?q={domain}"
    headers = {"User-Agent": USER_AGENT, "Accept": "application/json"}

    for intento in range(1, MAX_RETRIES + 1):
        try:
            if verbose: print(f"[~] Intento {intento}/{MAX_RETRIES}: consultando API JSON -> {url_json}")
            resp = requests.get(url_json, headers=headers, timeout=JSON_TIMEOUT)
            if verbose: print(f"[~] HTTP {resp.status_code}")
            if resp.status_code == 200:
                data = resp.json()
                for entry in data:
                    name_value = entry.get("name_value", "")
                    issuer = entry.get("issuer_name", "").lower()
                    is_lets = "Sí" if "let's encrypt" in issuer or "lets encrypt" in issuer else "No"
                    if name_value:
                        for s in name_value.split("\n"):
                            s = s.strip().lower()
                            if s and "*" not in s and s.endswith(domain):
                                subdomains[s] = is_lets
                if verbose: print(f"[+] Subdominios únicos encontrados vía API JSON: {len(subdomains)}")
                json_success = True
                break
        except Exception as e:
            if verbose: print(f"[!] Error en JSON: {e}")
        time.sleep(REQUESTS_SLEEP)

    # --- Fallback a HTML si JSON falló ---
    if not json_success:
        if verbose: print("[!] API JSON no disponible. Usando scraping HTML.")
        html_subs = scrape_html(domain, verbose=verbose)
        # Todos los del fallback se marcan como NO por instrucción del usuario
        subdomains = {s: "No" for s in html_subs}

    # --- Procesamiento final ---
    rows = []
    if verbose: print("[~] Resolviendo IPs y comprobando accesibilidad...")

    for i, (sub, lets) in enumerate(subdomains.items(), start=1):
        if verbose: print(f"[{i}/{len(subdomains)}] {sub} ... ", end="")
        ip = resolve_ip(sub)
        accede = check_accessibility(sub)
        if verbose: print(f"(IP: {ip}, Accede: {accede}, LetsEncrypt: {lets})")
        rows.append([sub, ip, accede, lets])

    df = pd.DataFrame(rows, columns=["Subdominio", "IP", "Accede http(s)", "LetsEncrypt"])
    if verbose:
        print(f"[+] Proceso completado en {time.time() - start_total:.2f}s. Total subdominios: {len(df)}")
    return df


def scrape_html(domain, verbose=True):
    url_html = f"https://crt.sh/?q={domain}"
    headers = {"User-Agent": USER_AGENT, "Accept": "text/html,application/xhtml+xml"}
    try:
        if verbose: print(f"[~] Solicitando HTML: {url_html}")
        resp = requests.get(url_html, headers=headers, timeout=HTML_TIMEOUT)
        if resp.status_code != 200:
            if verbose: print(f"[!] HTTP {resp.status_code}")
            return set()
        soup = BeautifulSoup(resp.text, "html.parser")
        subdomains = set()
        for td in soup.find_all("td"):
            txt = td.get_text(separator="\n").strip().lower()
            if domain in txt:
                for line in txt.splitlines():
                    line = line.strip()
                    if line.endswith(domain) and "*" not in line:
                        subdomains.add(line)
        if verbose: print(f"[+] Subdominios encontrados vía HTML: {len(subdomains)}")
        return subdomains
    except Exception as e:
        if verbose: print(f"[!] Error HTML: {e}")
        return set()


def resolve_ip(subdomain):
    try:
        return socket.gethostbyname(subdomain)
    except:
        return "No resuelve"


def check_accessibility(subdomain):
    headers = {"User-Agent": USER_AGENT}
    for scheme in ["https", "http"]:
        try:
            r = requests.get(f"{scheme}://{subdomain}", headers=headers, timeout=(4, 5))
            return r.status_code
        except:
            continue
    return "No"


# FUNCIONES PARA DOCX
def super_Tabla1(domain, df, document):
    p = document.add_heading(f"Informe de Subdominios - {domain}", level=1)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    t = document.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    hdr[0].text = "Nombre"
    hdr[1].text = "Severidad"
    hdr[2].text = "Descripción"
    r = t.add_row().cells
    r[0].text = "Enumeración de subdominios"
    r[1].text = "Baja"
    r[2].text = "Listado de subdominios identificados"

    document.add_paragraph()
    m = document.add_table(rows=1, cols=1)
    m.style = "Table Grid"
    m.rows[0].cells[0].text = "Recomendación: Revisar y controlar subdominios expuestos públicamente."

    document.add_paragraph()
    create_doc(domain, df, document)


def create_doc(domain, df, document):
    document.add_heading("Evidencia: Subdominios descubiertos", level=2)
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Subdominio"
    hdr[1].text = "IP"
    hdr[2].text = "Accede http(s)"
    hdr[3].text = "LetsEncrypt"

    for _, row in df.iterrows():
        cells = table.add_row().cells
        cells[0].text = str(row["Subdominio"])
        cells[1].text = str(row["IP"])
        cells[2].text = str(row["Accede http(s)"])
        cells[3].text = str(row["LetsEncrypt"])

    document.add_paragraph()
    document.add_heading("FUENTES", level=2)
    document.add_paragraph(f"https://crt.sh/json?q={domain}")


if __name__ == "__main__":
    domain = input("Ingresa el dominio principal: ").strip()
    df = fetch_data(domain, verbose=True)
    print(df)
