from docx import Document
import virustotal_subdomains as vt
import crtsh_subdomains as crt
import shodan_scrapping as sp
import severity_NIST as nist
#import ssl_tls_capture  # Importar ssl_tls para ejecutar el análisis SSL/TLS
import pandas as pd
from datetime import datetime
from docx.shared import Pt
import vulnerabilidades_to_doc as vulndoc


from infodns_to_doc import (
    agregar_tabla_analisis_perimetral,
    consulta_whois,
    consulta_dns_resuelta,
    consulta_dns
)


def combinar(dfvt, dfcrt):
    df = dfvt.merge(dfcrt, on="Subdominio", how="outer", suffixes=("_vt", "_crtsh"))
    df["IP"] = df[["IP_vt", "IP_crtsh"]].fillna("").agg("\n".join, axis=1).str.strip()
    df["Accede http(s)"] = df["Accede http(s)_vt"].combine_first(df["Accede http(s)_crtsh"])
    df = df[["Subdominio", "IP", "Accede http(s)"]]

    return df


def main(domain):

    # =========================
    # RECOLECCIÓN DNS
    # =========================
    whois_data = consulta_whois(domain)
    ns = consulta_dns_resuelta(domain, 'NS')
    a = consulta_dns_resuelta(domain, 'A')
    mx = consulta_dns_resuelta(domain, 'MX')
    txt = consulta_dns(domain, 'TXT')

    # =========================
    # SUBDOMINIOS
    # =========================
    dfvt = vt.fetch_data(domain)
    print(dfvt)
    
    dfcrt = crt.fetch_data(domain)
    print(dfcrt)
    
    df = combinar(dfvt, dfcrt)
    print(df)

    # =========================
    # SHODAN / PUERTOS / CVEs
    # =========================
    ip_df = sp.get_ips(df)
    puertos_df, allcve = sp.sh_scrapping(ip_df)

    # Puntajes CVSS desde NIST
    for ip in allcve.keys():
        allcve[ip] = nist.get_cves_base_scores(allcve[ip])

    print(ip_df)
    print(allcve)
    print(puertos_df)

    # =========================
    # DOCUMENTO WORD
    # =========================
    document = Document()

    # Fuente base
    style = document.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(9)

    # =========================
    # ANÁLISIS PERIMETRAL
    # =========================
    agregar_tabla_analisis_perimetral(
        document,
        domain,
        whois_data,
        ns,
        a,
        mx,
        txt
    )

    document.add_paragraph()  # 👈 separación

    # =========================
    # SUBDOMINIOS
    # =========================
    vt.super_Tabla1(domain, df, document)

    document.add_paragraph()  # 👈 separación

    # =========================
    # PUERTOS
    # =========================
    sp.super_Tabla2(puertos_df, document)

    document.add_paragraph()  # 👈 separación

    # =========================
    # VULNERABILIDADES
    # =========================
    for ip, cves in allcve.items():
        vulndoc.agregar_tabla_vulnerabilidades(document, ip, cves)
        document.add_paragraph()  # 👈 separación entre IPs

    # =========================
    # GUARDAR
    # =========================
    now = datetime.date(datetime.now())
    document.save(f'{domain}_{now}.docx')


if __name__ == "__main__":
    domains = input("Ingresa los dominios: ")
    domains = domains.split(' ')
    for domain in domains:
        main(domain)