from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from webdriver_manager.chrome import ChromeDriverManager
from docx import Document
from docx.shared import Inches
import os
import time

def ssl_test(domain, doc):
    options = Options()
    options.headless = True
    driver = webdriver.Chrome(service=Service(ChromeDriverManager().install()), options=options)
    driver.maximize_window()

    url = f'https://www.ssllabs.com/ssltest/analyze.html?d={domain}'
    driver.get(url)

    wait = WebDriverWait(driver, 120)
    try:
        elements = wait.until(EC.presence_of_all_elements_located((By.CLASS_NAME, 'reportSection')))
        
        # Agregar título del dominio en el documento
        doc.add_heading(f'Análisis SSL para {domain}', level=1)

        for i, element in enumerate(elements[:2]):
            driver.execute_script("arguments[0].scrollIntoView();", element)
            time.sleep(1)

            screenshot_name = f"{domain}_element_{i + 1}.png"
            element.screenshot(screenshot_name)
            print(f"[{domain}] Captura de pantalla guardada como '{screenshot_name}'")

            # Insertar imagen en el documento
            #doc.add_paragraph(f"Captura {i + 1} para {domain}:")
            doc.add_picture(screenshot_name, width=Inches(5))

            # Eliminar la imagen después de insertarla en el documento
            os.remove(screenshot_name)

    except Exception as e:
        print(f"[{domain}] Error: {e}")
        doc.add_paragraph(f"Error al analizar {domain}: {e}")

    finally:
        driver.quit()

if __name__ == "__main__":
    file_path = os.path.join(os.path.dirname(__file__), "dominios.txt")
    print(f"Buscando el archivo en: {file_path}")

    try:
        with open(file_path, "r") as file:
            domains = [line.strip() for line in file if line.strip()]

        # Crear documento de Word
        doc = Document()
        doc.add_heading("Reporte de Análisis SSL", level=0)

        for domain in domains:
            print(f"Iniciando análisis para: {domain}")
            ssl_test(domain, doc)

        # Guardar el documento al finalizar
        doc.save("ssl_reports.docx")
        print("\n📄 Reporte generado: ssl_reports.docx")

    except FileNotFoundError:
        print(f"Archivo no encontrado: {file_path}")
    except Exception as e:
        print(f"Error general: {e}")
