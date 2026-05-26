from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import whois
import dns.resolver
from ipwhois import IPWhois


# =========================
# UTILIDADES WORD
# =========================
def set_cell_background(cell, color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), color)
    tcPr.append(shd)


def format_text(cell, text, fuente, tamano):
    cell.text = text
    for p in cell.paragraphs:
        for run in p.runs:
            run.font.size = tamano
            run.font.name = fuente


# =========================
# FUNCIONES DNS / WHOIS
# =========================
def consulta_whois(dominio):
    try:
        return whois.whois(dominio)
    except:
        return None


def consulta_dns(dominio, tipo='A'):
    resultados = []
    try:
        respuestas = dns.resolver.resolve(dominio, tipo)
        for rdata in respuestas:
            resultados.append(str(rdata).strip())
    except:
        pass
    return resultados


def consulta_dns_resuelta(dominio, tipo='A'):
    registros = consulta_dns(dominio, tipo)
    resultado = {}

    if tipo == 'A':
        resultado[dominio] = registros
    else:
        for r in registros:
            resultado[r] = consulta_dns(r, 'A')

    return resultado


def consulta_whois_ip(ip):
    try:
        obj = IPWhois(ip)
        return obj.lookup_rdap()
    except:
        return None


# =========================
# FORMATEADORES
# =========================
def format_whois(info):
    if not info:
        return "Sin información WHOIS disponible"

    texto = []
    campos = ['domain_name', 'registrar', 'creation_date', 'expiration_date']

    for campo in campos:
        valor = info.get(campo)
        if valor:
            texto.append(f"{campo}: {valor}")

    return "\n".join(texto)


def format_dict_registros(data):
    salida = []
    for k, v in data.items():
        if v:
            salida.append(f"{k}: {', '.join(v)}")
        else:
            salida.append(f"{k}: sin resolución")
    return "\n".join(salida)


def format_lista(lista):
    return "\n".join(lista) if lista else "Sin registros"


# =========================
# FUNCIÓN PRINCIPAL (INTEGRABLE)
# =========================
def agregar_tabla_analisis_perimetral(document, dominio, whois_data, ns, a, mx, txt):

    fuente = "Calibri (Cuerpo)"
    tamano = Pt(9)
    header_dark = "666666"

    # =========================
    # TABLA PRINCIPAL
    # =========================
    table = document.add_table(rows=7, cols=2)
    table.style = 'Table Grid'

    merged = table.rows[0].cells[0].merge(table.rows[0].cells[1])
    set_cell_background(merged, "0070C0")

    header_data = [
        ("NOMBRE", "Análisis Perimetral"),
        ("SEVERIDAD", ""),
        ("ACCESO", dominio),
        ("RIESGO", "")
    ]

    for i, (k, v) in enumerate(header_data, start=1):
        set_cell_background(table.rows[i].cells[0], "D9D9D9")
        format_text(table.rows[i].cells[0], k, fuente, tamano)
        format_text(table.rows[i].cells[1], v, fuente, tamano)

    # =========================
    # HEADERS DESCRIPCIÓN / IMPACTO
    # =========================
    for i, txt_header in enumerate(["DESCRIPCIÓN", "IMPACTO"]):
        cell = table.rows[5].cells[i]
        set_cell_background(cell, header_dark)
        p = cell.paragraphs[0]
        run = p.add_run(txt_header)
        run.bold = True
        run.font.size = tamano
        run.font.name = fuente
        run.font.color.rgb = RGBColor(255, 255, 255)

    # =========================
    # TEXTO AUTOMÁTICO
    # =========================
    descripcion = (
        "Se realizó un análisis perimetral del dominio indicado, identificando registros DNS públicos asociados, incluyendo registros NS, A, MX y TXT.\n\n"
        "Esta información permite obtener visibilidad sobre la infraestructura expuesta a internet, así como identificar posibles superficies de ataque derivadas de configuraciones, servicios publicados y dependencias externas."
    )

    impacto = (
        "La exposición de información DNS puede facilitar tareas de reconocimiento por parte de un atacante, permitiendo identificar servicios accesibles, proveedores utilizados y posibles vectores de ataque.\n\n"
        "Configuraciones incorrectas o servicios asociados a estos registros podrían derivar en accesos no autorizados, uso de software desactualizado o explotación de vulnerabilidades conocidas."
    )

    format_text(table.rows[6].cells[0], descripcion, fuente, tamano)
    format_text(table.rows[6].cells[1], impacto, fuente, tamano)

    # =========================
    # BLOQUES DNS
    # =========================
    bloques = [
        ("REGISTROS DE DOMINIO", format_whois(whois_data)),
        ("REGISTROS NS", format_dict_registros(ns)),
        ("REGISTROS TIPO A", format_dict_registros(a)),
        ("REGISTROS MX", format_dict_registros(mx)),
        ("REGISTROS TXT", format_lista(txt))
    ]

    for titulo, contenido in bloques:
        t = document.add_table(rows=2, cols=1)
        t.style = 'Table Grid'

        set_cell_background(t.rows[0].cells[0], header_dark)

        p = t.rows[0].cells[0].paragraphs[0]
        run = p.add_run(titulo)
        run.bold = True
        run.font.size = tamano
        run.font.name = fuente
        run.font.color.rgb = RGBColor(255, 255, 255)

        format_text(t.rows[1].cells[0], contenido, fuente, tamano)